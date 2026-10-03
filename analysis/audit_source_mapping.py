"""Trace each model-input row back to an annual source workbook column."""

from __future__ import annotations

import csv
import re
from collections import Counter
from decimal import Decimal
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
FEATURE_ROWS = [*range(6, 36), 37, 38, 39, 40]
DEPTH = {"Forecourt Short Balls", "Midcourt Balls", "Backcourt Deep Balls"}


def number(value) -> float:
    if value is None:
        return float("nan")
    if isinstance(value, str):
        cleaned = re.sub(r"\s+", "", value.replace("、", ""))
        if cleaned in "零一二三四五六七八九":
            return float("零一二三四五六七八九".index(cleaned))
        return float(Decimal(cleaned))
    return float(value)


def main() -> None:
    mapping = pd.read_csv(DATA / "reference/match_mapping.csv")
    train = pd.read_csv(DATA / "processed/train_grouped.csv")
    test = pd.read_csv(DATA / "processed/test_grouped.csv")
    features = list(train.columns[1:35])
    assert len(mapping) == len(train) + len(test) == 239
    assert len(features) == len(FEATURE_ROWS) == 34
    workbooks = {}
    differences = []
    for item in mapping.itertuples(index=False):
        source = DATA / "raw" / item.source_file
        if source not in workbooks:
            workbooks[source] = load_workbook(source, read_only=False, data_only=True)
        sheet = workbooks[source][item.source_sheet]
        column = column_index_from_string(item.source_column)
        table = train if item.split == "train" else test
        row = table.iloc[int(item.csv_data_row) - 1]
        assert row.match_id == item.match_id
        assert int(row["Set Result"]) == int(item.label)
        assert int(row.year) == int(item.year)
        assert int(row.set_within_match) == int(item.set_within_match)
        source_rows = [*range(5, 35), 36, 37, 38, 39] if int(item.year) == 2021 else FEATURE_ROWS
        for feature, source_row in zip(features, source_rows):
            recorded = number(sheet.cell(source_row, column).value)
            used = float(row[feature])
            difference = used - recorded
            if not np.isfinite(recorded) or abs(difference) > 1e-9:
                differences.append({"split": item.split, "year": item.year,
                                    "match_id": item.match_id, "set_within_match": item.set_within_match,
                                    "feature": feature, "source_cell": f"{item.source_column}{source_row}",
                                    "source_value": recorded, "model_value": used, "difference": difference})
    for workbook in workbooks.values():
        workbook.close()
    output = ROOT / "outputs/source_value_differences.csv"
    output.parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(differences).to_csv(output, index=False)
    print(f"Compared {239 * 34} source values; differences: {len(differences)}")
    print(Counter((r["year"], r["feature"]) for r in differences).most_common(25))


if __name__ == "__main__":
    main()
