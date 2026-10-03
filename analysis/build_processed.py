"""Rebuild the model tables from annual workbooks and the set-to-source map."""

from __future__ import annotations

from decimal import Decimal, ROUND_FLOOR, ROUND_HALF_UP
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from audit_source_mapping import ROOT, DATA, FEATURE_ROWS, number


RATE_ROUNDING_CELLS = {(2019, "G22"), (2019, "I22"), (2019, "O22"), (2024, "J21")}
DEPTH_NORMALIZATION_SETS = {
    ("2020_AO_F_Djokovic", 3),
    ("2025_AO_QF_Djokovic", 1),
    ("2025_USO_R128_Djokovic", 2),
}


def normalize_depth(values: list[float]) -> list[float]:
    raw = [Decimal(str(x)) for x in values]
    shares = [v * 100 / sum(raw) for v in raw]
    counts = [int(v.to_integral_value(rounding=ROUND_FLOOR)) for v in shares]
    remaining = 100 - sum(counts)
    order = sorted(range(3), key=lambda i: (-(shares[i] - counts[i]), i))
    for i in order[:remaining]:
        counts[i] += 1
    return [n / 100 for n in counts]


def main() -> None:
    mapping = pd.read_csv(DATA / "reference/match_mapping.csv")
    expected = {
        "train": pd.read_csv(DATA / "processed/train_grouped.csv"),
        "test": pd.read_csv(DATA / "processed/test_grouped.csv"),
    }
    features = list(expected["train"].columns[1:35])
    workbooks = {}
    rows = {"train": [], "test": []}
    for item in mapping.itertuples(index=False):
        source = DATA / "raw" / item.source_file
        if source not in workbooks:
            workbooks[source] = load_workbook(source, read_only=False, data_only=True)
        sheet = workbooks[source][item.source_sheet]
        column = column_index_from_string(item.source_column)
        source_rows = [*range(5, 35), 36, 37, 38, 39] if int(item.year) == 2021 else FEATURE_ROWS
        values = []
        for row_number in source_rows:
            value = number(sheet.cell(row_number, column).value)
            assert np.isfinite(value)
            if (int(item.year), f"{item.source_column}{row_number}") in RATE_ROUNDING_CELLS:
                value = float(Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
            values.append(value)
        if (item.match_id, int(item.set_within_match)) in DEPTH_NORMALIZATION_SETS:
            values[-3:] = normalize_depth(values[-3:])
        assert len(values) == 34
        row = {"Set Result": int(item.label), **dict(zip(features, values)),
               "match_id": item.match_id, "year": int(item.year), "event": item.event,
               "round": item.round, "set_within_match": int(item.set_within_match),
               "original_data_row": int(item.csv_data_row), "cv_group": item.match_id,
               "confirmed_speed_correction": int(int(item.year) == 2024 and item.source_column == "Y")}
        rows[item.split].append(row)
    for workbook in workbooks.values():
        workbook.close()
    output = ROOT / "outputs"
    output.mkdir(parents=True, exist_ok=True)
    for split in ("train", "test"):
        rebuilt = pd.DataFrame(rows[split], columns=expected[split].columns)
        reference = expected[split]
        assert rebuilt.shape == reference.shape
        assert np.allclose(rebuilt.iloc[:, :35].to_numpy(float), reference.iloc[:, :35].to_numpy(float), atol=1e-10)
        for column in reference.columns[35:]:
            assert np.array_equal(rebuilt[column].to_numpy(str), reference[column].to_numpy(str)), column
        rebuilt.to_csv(output / f"rebuilt_{split}_from_workbooks.csv", index=False)
    print("Rebuilt all 239 set rows and 34 candidate indicators from the six workbooks")


if __name__ == "__main__":
    main()
